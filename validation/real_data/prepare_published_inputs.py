"""Download deposited inputs; no orthology, solving or rendering here."""
import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path
import urllib.request
import shutil
import subprocess
import zipfile

PLANARIAN_URL = ('https://media.springernature.com/original/springer-static/esm/'
    'art%3A10.1038%2Fs41467-024-52380-9/MediaObjects/41467_2024_52380_MOESM4_ESM.xlsx')
PLANARIAN_SHA256 = '7d26221e5f29dc838c5d37267864f5e4093f6f782888734495f20b57a78a9953'
PLANARIAN_SPECIES = ['schLug1', 'schNov1', 'schPol2', 'schMedS3h2', 'schMedS3h1']
LEPTIDEA = {
    'LsinapisSpaM': 'GCA_949711715.1', 'LsinapisSweM': 'GCA_949711785.1',
    'LjuvernicaM': 'GCA_949711755.1', 'LrealiM': 'GCA_949710795.1',
}


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''): h.update(chunk)
    return h.hexdigest()


def download(url, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists(): return
    tmp = path.with_name(path.name + '.partial')
    print(f'DOWNLOADING {url}', flush=True)
    try:
        if shutil.which('curl'):
            subprocess.run([
                'curl', '--fail', '--location', '--show-error',
                '--retry', '3', '--retry-all-errors', '--retry-delay', '2',
                '--connect-timeout', '60', '--max-time', '600',
                '--output', str(tmp), url,
            ], check=True)
        else:
            request = urllib.request.Request(url, headers={'User-Agent': 'SynTangle-public-input-preparation/1'})
            with urllib.request.urlopen(request, timeout=180) as r, tmp.open('wb') as f:
                while chunk := r.read(1024 * 1024): f.write(chunk)
        if not tmp.stat().st_size:
            raise ValueError(f'Empty download: {url}')
        tmp.replace(path)
    finally: tmp.unlink(missing_ok=True)


def planarians(out):
    import openpyxl
    from syntangle.fixtures import fixture_from_dict
    from prepare_genespace_blocks import natural
    source = out / 'Supplementary_Data_1-15.xlsx'
    download(PLANARIAN_URL, source)
    if digest(source) != PLANARIAN_SHA256:
        raise ValueError('Planarian workbook differs from inspected source; inspect before importing')
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    iterator = workbook['Supplementary Data 14'].iter_rows(values_only=True)
    headers = next(iterator)
    required = {'genome1', 'genome2', 'chr1', 'chr2', 'blkID', 'orient',
                'startBp1', 'endBp1', 'startBp2', 'endBp2'}
    if not required.issubset(headers): raise ValueError('Unexpected block columns')
    rows = [dict(zip(headers, row)) for row in iterator]
    with (out/'published_block_coordinates.tsv').open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter='\t')
        writer.writeheader(); writer.writerows(rows)
    gene_iterator = workbook['Supplementary Data 13'].iter_rows(values_only=True)
    gene_headers = next(gene_iterator)
    if not {'chr', 'start', 'end', 'genome', 'og'}.issubset(gene_headers):
        raise ValueError('Unexpected published gene-coordinate columns')
    with (out/'published_gene_coordinates.tsv').open('w', newline='') as f:
        writer = csv.writer(f, delimiter='\t'); writer.writerow(gene_headers)
        writer.writerows(gene_iterator)
    # Retain the author script's five-row chain. Only explicit chromosome rows
    # enter this chromosome-scale projection; every exclusion is catalogued.
    chromosome_blocks = {sp: defaultdict(list) for sp in PLANARIAN_SPECIES}
    edges, exclusions, pairs = [], [], Counter()
    for line, row in enumerate(rows, 2):
        a, b = row['genome1'], row['genome2']
        side_order = None
        for left, right in zip(PLANARIAN_SPECIES, PLANARIAN_SPECIES[1:]):
            if (a, b) == (left, right): side_order = (1, 2); break
            if (a, b) == (right, left): side_order = (2, 1); break
        if side_order is None:
            exclusions.append({'source_row': line, 'reason': 'outside_adjacent_chain'}); continue
        if not all(str(row[f'chr{i}']).startswith('chr') for i in (1, 2)):
            exclusions.append({'source_row': line, 'reason': 'unplaced_scaffold'}); continue
        hid = f'published_row_{line}'
        for sp, i in zip((left, right), side_order):
            start, end = sorted((int(row[f'startBp{i}']), int(row[f'endBp{i}'])))
            if start < 0 or start == end: raise ValueError(f'Invalid span in source row {line}')
            chromosome_blocks[sp][row[f'chr{i}']].append(dict(
                occurrence_id=hid+'_'+sp, homology_id=hid, start=start, end=end, strand='?'))
        edges.append(dict(homology_id=hid, source_row=line, source_block_id=row['blkID'],
                          source_orientation=row['orient'], original_genomes=[a, b]))
        pairs[left, right] += 1
    workbook.close()
    if len(pairs) != len(PLANARIAN_SPECIES)-1: raise ValueError('Missing adjacent pair')
    # Two source directions would duplicate evidence, so stop rather than merge.
    directions = defaultdict(set)
    for edge in edges:
        a, b = edge['original_genomes']; directions[frozenset((a, b))].add((a, b))
    if any(len(x) != 1 for x in directions.values()): raise ValueError('Reciprocal rows need explicit reconciliation')
    fixture = dict(fixture_version=1, id='planarians_published_blocks',
        title='Schmidtea: deposited GENESPACE blocks', purpose='Published chromosome block projection',
        orientation_constraints=[], species=[])
    for sp in PLANARIAN_SPECIES:
        fixture['species'].append(dict(id=sp, chromosomes=[dict(id=chrom,
            length=max(b['end'] for b in blocks), display_rank=rank+1, blocks=blocks)
            for rank, (chrom, blocks) in enumerate(sorted(chromosome_blocks[sp].items(), key=lambda x: natural(x[0])))]))
    fixture_from_dict(fixture)
    linked_chromosomes = sum(len(s['chromosomes']) for s in fixture['species'])
    if len(edges) != 585 or linked_chromosomes != 19:
        raise ValueError('Chromosome projection differs from inspected source counts')
    provenance = dict(paper='https://doi.org/10.1038/s41467-024-52380-9', source_url=PLANARIAN_URL,
        sha256=PLANARIAN_SHA256, sheet='Supplementary Data 14', source_rows=len(rows),
        retained_links=len(edges), edges=edges, excluded_rows=exclusions,
        pair_audit=[dict(pair=list(pair), retained=n) for pair, n in pairs.items()],
        species_row_order='Author genespace_plotting.R: reversed five-assembly list',
        chromosome_extent='Maximum retained block endpoint; NOT assembly chromosome lengths',
        display_order='Natural chromosome-name order; NOT recovered published chromosome order/flips',
        objective='Unweighted adjacent-layer block-midpoint crossings; NOT original gene-anchor objective',
        exact_published_display=False,
        scope='Chromosome-only projection; all retained links conserved, exclusions explicitly listed')
    (out/'fixture.json').write_text(json.dumps(fixture, indent=2)+'\n')
    (out/'provenance.json').write_text(json.dumps(provenance, indent=2)+'\n')
    print(f'PREPARED planarians: {len(edges)} links; {sum(len(s["chromosomes"]) for s in fixture["species"])} linked chromosomes', flush=True)


def leptidea(out):
    records = []
    for name, accession in LEPTIDEA.items():
        base = f'https://api.ncbi.nlm.nih.gov/datasets/v2/genome/accession/{accession}'
        report = out / name / 'dataset_report.json'
        download(base+'/dataset_report', report)
        reports = json.loads(report.read_text())['reports']
        if len(reports) != 1 or reports[0]['accession'] != accession:
            raise ValueError('Unexpected assembly accession')
        info = reports[0]['assembly_info']
        if info['bioproject_accession'] != 'PRJEB58697' or info['assembly_name'] != name:
            raise ValueError('Assembly does not match the published project/name')
        url = base+'/download?include_annotation_type=GENOME_FASTA&include_annotation_type=SEQUENCE_REPORT'
        archive = out / name / (accession+'.zip')
        print('DOWNLOADING', name, accession, flush=True)
        download(url, archive)
        with zipfile.ZipFile(archive) as z:
            if z.testzip() is not None: raise ValueError('Damaged assembly archive')
            fasta = [p for p in z.namelist() if p.endswith('.fna')]
            if not fasta: raise ValueError('Downloaded assembly contains no FASTA')
        records.append(dict(name=name, accession=accession, url=url,
                            sha256=digest(archive), archive=str(archive), fasta_members=fasta))
    (out/'provenance.json').write_text(json.dumps(dict(
        paper='https://doi.org/10.1007/s10577-023-09713-z', assemblies=records,
        status='assemblies_downloaded_NOT_plot_ready',
        remaining=['Verify original MAKER annotations and proteins',
                   'Pin original Bombyx and Melitaea reference versions',
                   'Rerun published BLASTP/MCScanX pipeline: maximum gene gap 10',
                   'Recover SynVisio row order and chromosome-size order'],
        warning='No alternative DNA-alignment evidence substituted for published gene-based links'), indent=2)+'\n')
    print('DOWNLOADED Leptidea assemblies; annotations and published collinearity still required', flush=True)


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--dataset', choices=['planarians', 'leptidea'], required=True)
    p.add_argument('--output', type=Path, required=True); args = p.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    globals()[args.dataset](args.output)
